"""Cuts a blinded review workbook down to a hand-scorable stratified sample.

`build_review_sheet.py` emits every claim the judge produced, which was 1,697 for the
40-row sample. That is not scorable by one annotator. This draws a fixed-seed subset
balanced across (model, context_variant), spread evenly over repositories inside each
cell so no single verbose summary dominates, and writes a workbook with the same
structure: a visible blinded Review sheet, a hidden Answer Key, and the original
`row_id` values preserved so `scripts.score_human_validation` still joins correctly.

Run from backend/:

    python -m scripts.subsample_review_sheet
    python -m scripts.subsample_review_sheet --per-cell 20 --seed 7
"""

from __future__ import annotations

import argparse
import os
import random
import sys
from collections import defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_BACKEND_DIR = os.path.dirname(_SCRIPT_DIR)
_DEFAULT_RESULTS_DIR = os.path.join(_BACKEND_DIR, "evaluation_results")
_DEFAULT_IN = os.path.join(_DEFAULT_RESULTS_DIR, "human_review.xlsx")
_DEFAULT_OUT = os.path.join(_DEFAULT_RESULTS_DIR, "human_review_sample.xlsx")

INSTRUCTIONS = [
    ("How to score this sheet", True),
    ("", False),
    ("You are checking whether each claim about a repository is supported by that "
     "repository's real structure. You are not judging whether the claim is "
     "well written, useful, or complete.", False),
    ("", False),
    ("1. Work on the 'Review' tab. Fill in the 'human_verdict' column only. "
     "Leave every other column alone.", False),
    ("2. Each block of rows shares one repository. The framework, known_dependencies "
     "and known_endpoints columns are filled in on the first row of each block and "
     "apply to every row beneath it until the repository name changes.", False),
    ("3. For each claim, pick one value from the dropdown:", False),
    ("      Supported    - the repository facts shown for that repository confirm the "
     "claim, or the claim is a fair restatement of them.", False),
    ("      Unsupported  - the facts contradict the claim, or the claim asserts "
     "something specific that the facts do not contain at all.", False),
    ("      Unsure       - you genuinely cannot tell from the facts given. Use this "
     "sparingly; it is reported separately and does not count against anyone.", False),
    ("", False),
    ("Rules that matter for the numbers:", True),
    ("  - Judge only against the facts in the row's repository block. Do not open "
     "GitHub, do not guess from the repository name, and do not reward a claim for "
     "being plausible.", False),
    ("  - A vague claim ('this is a web application') is Supported. Only specific, "
     "checkable claims can be Unsupported.", False),
    ("  - A claim naming a dependency, endpoint or framework that is absent from the "
     "block is Unsupported, even if such a thing would be normal for this kind of "
     "project.", False),
    ("  - Score rows in the order they appear. Do not sort or delete rows.", False),
    ("  - Do not unhide or look at the 'Answer Key' tab. It holds the model's own "
     "verdicts, and reading it first invalidates the agreement measurement this "
     "sheet exists to produce.", False),
    ("", False),
    ("Save the file when done and hand it back unchanged apart from the "
     "human_verdict column.", False),
]


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in-sheet", default=_DEFAULT_IN)
    ap.add_argument("--out", default=_DEFAULT_OUT)
    ap.add_argument("--per-cell", type=int, default=28,
                    help="claims per (model, representation) cell; 9 cells, so 28 gives 252")
    ap.add_argument("--seed", type=int, default=42)
    a = ap.parse_args()

    if not os.path.exists(a.in_sheet):
        print(f"FATAL: {a.in_sheet} not found. Run scripts.build_review_sheet first.",
              file=sys.stderr)
        return 2

    wb = load_workbook(a.in_sheet)
    for name in ("Review", "Answer Key"):
        if name not in wb.sheetnames:
            print(f"FATAL: {a.in_sheet} has no '{name}' worksheet.", file=sys.stderr)
            return 2
    ws_r, ws_k = wb["Review"], wb["Answer Key"]

    hdr_r = [c.value for c in ws_r[1]]
    hdr_k = [c.value for c in ws_k[1]]
    if "row_id" not in hdr_r or "row_id" not in hdr_k:
        print("FATAL: both worksheets need a 'row_id' column. Regenerate with "
              "scripts.build_review_sheet.", file=sys.stderr)
        return 2

    review = {}
    for row in ws_r.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(hdr_r, row))
        review[rec["row_id"]] = rec
    key = {}
    for row in ws_k.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(hdr_k, row))
        key[rec["row_id"]] = rec

    missing = sorted(set(review) - set(key))
    if missing:
        print(f"FATAL: {len(missing)} Review row(s) absent from the Answer Key.",
              file=sys.stderr)
        return 1

    # cell -> repo -> [row_id]
    cells: dict[tuple, dict[str, list]] = defaultdict(lambda: defaultdict(list))
    for rid, rec in key.items():
        cell = (rec["model"], rec["context_variant"])
        cells[cell][rec["repo_name"]].append(rid)

    rng = random.Random(a.seed)
    chosen: list = []
    print(f"claims available: {len(review)} across {len(cells)} (model, representation) cells")
    for cell in sorted(cells):
        repos = sorted(cells[cell])
        for r in repos:
            rng.shuffle(cells[cell][r])
        picked = []
        # round-robin over repositories so the cell's quota spreads across them
        i = 0
        while len(picked) < a.per_cell and any(cells[cell][r] for r in repos):
            r = repos[i % len(repos)]
            if cells[cell][r]:
                picked.append(cells[cell][r].pop())
            i += 1
        chosen.extend(picked)
        print(f"  {cell[0]:<22} {cell[1]:<18} {len(picked):3d} claims "
              f"from {len(set(key[p]['repo_name'] for p in picked))} repos")

    # Original row_id order keeps repositories grouped the way the reviewer expects.
    chosen.sort()
    print(f"sampled {len(chosen)} claims")

    out = Workbook()
    ws_i = out.active
    ws_i.title = "Instructions"
    ws_i.column_dimensions["A"].width = 118
    for i, (text, bold) in enumerate(INSTRUCTIONS, start=1):
        c = ws_i.cell(row=i, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if bold:
            c.font = Font(bold=True)

    ws_out = out.create_sheet("Review")
    ws_key = out.create_sheet("Answer Key")
    ws_key.sheet_state = "hidden"

    ws_out.append(hdr_r)
    ws_key.append(hdr_k)

    col = {n: i for i, n in enumerate(hdr_r, start=1)}
    verdict_letter = get_column_letter(col["human_verdict"])
    dv = DataValidation(type="list", formula1='"Supported,Unsupported,Unsure"',
                        allow_blank=True, showDropDown=False)
    ws_out.add_data_validation(dv)

    last_repo = None
    for i, rid in enumerate(chosen, start=2):
        rec = dict(review[rid])
        # Repeated repository facts are blanked as in the full sheet, but the first row
        # of each block must carry them, and sub-sampling changes which row that is.
        if rec["repo_name"] == last_repo:
            for f in ("framework", "known_dependencies", "known_endpoints"):
                rec[f] = ""
        else:
            src = key[rid]
            for f in ("framework", "known_dependencies", "known_endpoints"):
                rec[f] = src.get(f, "")
        last_repo = rec["repo_name"]
        ws_out.append([rec.get(h, "") for h in hdr_r])
        ws_key.append([key[rid].get(h, "") for h in hdr_k])
        dv.add(f"{verdict_letter}{i}")

    widths = {"row_id": 8, "repo_name": 26, "known_dependencies": 30,
              "known_endpoints": 34, "claim_text": 66, "human_verdict": 16}
    for name, w in widths.items():
        if name in col:
            ws_out.column_dimensions[get_column_letter(col[name])].width = w
    for r in range(2, len(chosen) + 2):
        for name in ("known_endpoints", "claim_text"):
            if name in col:
                ws_out.cell(row=r, column=col[name]).alignment = Alignment(wrap_text=True,
                                                                           vertical="top")
    for c in ws_out[1]:
        c.font = Font(bold=True)
    ws_out.freeze_panes = "A2"
    ws_out.auto_filter.ref = None

    os.makedirs(os.path.dirname(a.out) or ".", exist_ok=True)
    out.save(a.out)
    print(f"wrote {a.out}")
    print(f"reviewer fills column {verdict_letter} on the 'Review' tab; "
          f"score it with `python -m scripts.score_human_validation --sheet {a.out}`")
    return 0


if __name__ == "__main__":
    sys.exit(main())
